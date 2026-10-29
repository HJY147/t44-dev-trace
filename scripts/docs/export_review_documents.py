"""Export four private technical review DOCX files from the project's own Markdown.

Run only with the bundled Python. Historical M6/M7 continue their original operation;
the separately authorized FINAL-20261029 batch has one new create-four marker.
No product imports, requests, tests or generic document templates are used.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import stat
import shutil
from datetime import datetime, timezone
import subprocess
import sys
from pathlib import Path
from urllib.parse import quote
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

RUNTIME = Path('/Users/yang/.cache/codex-runtimes/codex-primary-runtime/dependencies')
REPORTS = [
    ('requirements-report.md', 'requirements-review.docx', 'T44软件需求规格说明书'),
    ('design-report.md', 'design-review.docx', 'T44概要设计与详细设计报告'),
    ('test-results-report.md', 'test-results-review.docx', 'T44测试结果技术报告'),
    ('ai-spec-guide.md', 'ai-spec-guide-review.docx', 'T44 AI可执行规格相关文件说明'),
]
ASSET_HASHES = {}


def font(run, size=None, bold=None):
    run.font.name = 'Songti SC'
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for key, val in [('ascii', 'Songti SC'), ('hAnsi', 'Songti SC'), ('eastAsia', 'Songti SC')]:
        fonts.set(qn('w:' + key), val)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def plain(text):
    return re.sub(r'[`*_]', '', text).strip()


def href(target, source, root, cutoff):
    if target.startswith(('https://', 'http://', 'mailto:')):
        return target
    path, sep, anchor = target.partition('#')
    resolved = (source.parent / path).resolve() if path else source
    relative = resolved.relative_to(root).as_posix()
    return ('https://github.com/HJY147/t44-dev-trace/blob/' + cutoff + '/' + quote(relative)
            + ('#' + anchor if sep else ''))


def link(paragraph, label, target, source, root, cutoff, size=None):
    item = OxmlElement('w:hyperlink')
    rid = paragraph.part.relate_to(href(target, source, root, cutoff),
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
        is_external=True)
    item.set(qn('r:id'), rid)
    r = OxmlElement('w:r')
    rp = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    for k, v in [('ascii', 'Songti SC'), ('hAnsi', 'Songti SC'), ('eastAsia', 'Songti SC')]:
        rf.set(qn('w:' + k), v)
    rp.append(rf)
    color = OxmlElement('w:color'); color.set(qn('w:val'), '1F4E79'); rp.append(color)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int((size or 11) * 2))); rp.append(sz)
    r.append(rp); t = OxmlElement('w:t'); t.text = label; r.append(t); item.append(r)
    paragraph._p.append(item)


def inline(paragraph, text, source, root, cutoff, size=None):
    # The source uses simple links, strong text and code, without nested inline markup.
    pattern = r'(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)'
    for part in re.split(pattern, text):
        if not part:
            continue
        m = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', part)
        if m:
            link(paragraph, m[1], m[2], source, root, cutoff, size)
        else:
            code = part.startswith('`') and part.endswith('`')
            strong = part.startswith('**') and part.endswith('**')
            run = paragraph.add_run(part[1:-1] if code else part[2:-2] if strong else part)
            font(run, size, strong or None)
            if code:
                run.font.name = 'Consolas'
                run.font.size = Pt(min(size or 10.5, 10.5))
    ppr = paragraph._p.get_or_add_pPr()
    wrap = OxmlElement('w:wordWrap'); wrap.set(qn('w:val'), '0'); ppr.append(wrap)


def paragraph(doc, text, source, root, cutoff, *, style=None):
    p = doc.add_paragraph(style=style)
    inline(p, text, source, root, cutoff)
    p.paragraph_format.widow_control = True
    if len(text) <= 260:
        p.paragraph_format.keep_together = True
    return p


def configure(doc, title):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(.75), Inches(.7)
    sec.left_margin, sec.right_margin = Inches(.75), Inches(.75)
    sec.header_distance, sec.footer_distance = Inches(.3), Inches(.3)
    sec.different_first_page_header_footer = True
    for name, size in [('Normal', 11), ('Title', 26), ('Subtitle', 14),
                       ('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11.5)]:
        st = doc.styles[name]
        st.font.name = 'Songti SC'; st.font.size = Pt(size); st.font.color.rgb = RGBColor(0, 0, 0)
        st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Songti SC')
        st.paragraph_format.space_after = Pt(7 if name == 'Normal' else 9)
        st.paragraph_format.line_spacing = Pt(16 if name == 'Normal' else size * 1.25)
        ppr = st._element.get_or_add_pPr()
        snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0'); ppr.append(snap)
        if name.startswith('Heading'):
            st.paragraph_format.space_before = Pt(14)
            st.paragraph_format.keep_with_next = True
    h = sec.header.paragraphs[0]; h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(h.add_run(title), 8.5)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run('第 '), 9)
    f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), 'PAGE'); footer._p.append(f)
    font(footer.add_run(' 页'), 9)
    for style in doc.styles:
        for border in style._element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
        if style.style_id == 'Title':
            for underline in style._element.xpath('.//w:u'):
                underline.getparent().remove(underline)
    props = doc.core_properties
    props.title = title; props.subject = '截止版本技术审阅材料 未通过最终验收'
    props.author = 'AI文档岗整理'; props.last_modified_by = 'AI文档岗整理'
    props.comments = '学生研究正文不在本次四份技术材料内'


def cover(doc, title, info, cutoff):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(38)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run(info.get('course', '软件设计与开发实验')), 21, True)
    p.paragraph_format.space_after = Pt(38)
    p = doc.add_paragraph(title, 'Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs: font(r, 26, True)
    p.paragraph_format.space_after = Pt(20)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run('技术报告'), 15, True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run('验收状态以测试结果报告为准'), 12)
    p.paragraph_format.space_after = Pt(38)
    for key, label in [('class_name', '专业班级'), ('student_id', '学号'),
                       ('student_name', '姓名'), ('instructor', '指导教师')]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(p.add_run(label + '：' + str(info.get(key, ''))), 12)
        p.paragraph_format.space_after = Pt(12)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run('日期：'), 12)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(p.add_run('文档汇总截止 ' + cutoff[:7] + '  测试证据按各自准确SHA'), 9)
    doc.add_page_break()


def table(doc, lines, source, root, cutoff):
    rows = [[x.strip() for x in line.strip().strip('|').split('|')] for line in lines]
    data = [rows[0]] + rows[2:]
    n = len(data[0]); t = doc.add_table(rows=0, cols=n); t.autofit = False
    width = (doc.sections[-1].page_width - doc.sections[-1].left_margin - doc.sections[-1].right_margin) / 914400
    if n >= 5:
        fractions = ([.09, .25, .28, .28, .1] if n == 5 else [1/n]*n)
    else:
        weights = [min(80, max(12, sum(len(plain(r[i])) for r in data)/len(data))) for i in range(n)]
        weights[0] = min(weights[0], 25)
        fractions = [v/sum(weights) for v in weights]
    for column, fraction in zip(t.columns, fractions): column.width = Inches(width*fraction)
    for index, values in enumerate(data):
        cells = t.add_row().cells
        no_split = OxmlElement('w:cantSplit')
        t.rows[-1]._tr.get_or_add_trPr().append(no_split)
        for i,c in enumerate(cells):
            c.width = Inches(width*fractions[i]); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = c._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge,val in [('top','75'),('bottom','75'),('left','85'),('right','85')]:
                node = OxmlElement('w:'+edge);node.set(qn('w:w'),val);node.set(qn('w:type'),'dxa');margins.append(node)
            tcpr.append(margins)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                b = OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            tcpr.append(borders)
            if index == 0:
                shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E8EEF3');tcpr.append(shading)
            para=c.paragraphs[0];para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=Pt(14)
            if i == 0 and len(plain(values[i])) < 20:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
            inline(para,values[i],source,root,cutoff,size=10 if n>=5 else 10.5)
            if index==0:
                for r in para.runs:r.bold=True
        if index == 0:
            repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(3)


def figure(doc, f, number, root):
    asset=root/f['artifacts']['png']['path']
    if hashlib.sha256(asset.read_bytes()).hexdigest()!=f['artifacts']['png']['sha256']:
        raise ValueError('Diagram asset checksum mismatch')
    w,h=Image.open(asset).size
    maximum_height = 6.0 if Path(f['mermaid']).stem.endswith('use-cases') else 7.1
    width=min(6.85,maximum_height*w/h);height=width*h/w
    caption=doc.add_paragraph();caption.paragraph_format.keep_with_next=True
    caption.alignment=WD_ALIGN_PARAGRAPH.CENTER
    labels={'context':'系统边界','use-cases':'角色和主要用例','states':'主要状态流转',
            'blocked':'阻塞与恢复','architecture':'模块与进程架构','entities':'核心实体关系',
            'transaction':'成功拒绝与重放时序','notifications':'扫描消息与降级时序'}
    stem=Path(f['mermaid']).stem
    name=next((v for k,v in labels.items() if stem.endswith(k)),stem)
    font(caption.add_run('图 '+str(number)+' '+name),10,True)
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.line_spacing = 1.0  # Inline drawings must set their own full line height.
    para.add_run().add_picture(str(asset),width=Inches(width),height=Inches(height))
    para.paragraph_format.keep_with_next=False
    para.paragraph_format.space_after=Pt(9)
    inline_shape=para._p.xpath('.//wp:docPr')[0]
    inline_shape.set('descr',name+' 复用M3已渲染图，非本轮重绘')


def image_figure(doc, source, root, target, title, number, *, restore_portrait=True):
    asset = (source.parent / target).resolve()
    relative = asset.relative_to(root).as_posix()
    if relative not in ASSET_HASHES or hashlib.sha256(asset.read_bytes()).hexdigest() != ASSET_HASHES[relative]:
        raise ValueError('Markdown figure is absent from the verified diagram manifests')
    w, h = Image.open(asset).size
    wide = w / h > 1.45 or asset.stem in ('m6-reason-projection', 'm6-quality-feedback')
    preceding_heading = None
    if wide:
        last = doc.paragraphs[-1]
        if last.style.name.startswith('Heading'):
            preceding_heading = (last.text, last.style)
            last._element.getparent().remove(last._element)
            # A table's spacer must not survive alone before the section break.
            while doc.paragraphs and not doc.paragraphs[-1].text and not doc.paragraphs[-1]._p.xpath('.//w:sectPr|.//w:drawing'):
                empty = doc.paragraphs[-1]._element
                empty.getparent().remove(empty)
        sec = doc.add_section(WD_SECTION.NEW_PAGE)
        boundary = doc.paragraphs[-1].paragraph_format
        boundary.space_before = boundary.space_after = Pt(0)
        boundary.line_spacing = Pt(1)
        boundary.keep_with_next = boundary.keep_together = False
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Inches(11), Inches(8.5)
        sec.different_first_page_header_footer = False
        if preceding_heading:
            doc.add_paragraph(preceding_heading[0], preceding_heading[1])
    maximum_width, maximum_height = (9.4, 6.0 if preceding_heading else 6.3) if wide else (6.85, 7.7)
    if asset.stem == 'requirements-report-03-states':
        maximum_height = 7.0
    width = min(maximum_width, maximum_height * w / h)
    height = width * h / w
    caption = doc.add_paragraph()
    caption.paragraph_format.keep_with_next = True
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(caption.add_run('图 ' + str(number) + ' ' + plain(title)), 10, True)
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.line_spacing = 1.0
    para.add_run().add_picture(str(asset), width=Inches(width), height=Inches(height))
    para.paragraph_format.space_after = Pt(5)
    para._p.xpath('.//wp:docPr')[0].set('descr', plain(title))
    if wide and restore_portrait:
        sec = doc.add_section(WD_SECTION.NEW_PAGE)
        boundary = doc.paragraphs[-1].paragraph_format
        boundary.space_before = boundary.space_after = Pt(0)
        boundary.line_spacing = Pt(1)
        boundary.keep_with_next = boundary.keep_together = False
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.different_first_page_header_footer = False


def body(doc, source, root, cutoff, figures):
    lines=source.read_text().splitlines(); i=0; fig=0; image_count=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('# '):i+=1;continue
        image = re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)', line)
        if image:
            image_count += 1
            j = i + 1
            while j < len(lines) and (not lines[j].strip() or re.match(r"^#{2,4}\s", lines[j].strip())):
                j += 1
            next_image = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", lines[j].strip()) if j < len(lines) else None
            next_wide = False
            if next_image:
                next_asset = (source.parent / next_image[2]).resolve()
                with Image.open(next_asset) as next_info:
                    next_wide = next_info.width / next_info.height > 1.45 or next_asset.stem in ('m6-reason-projection', 'm6-quality-feedback')
            image_figure(doc, source, root, image[2], image[1], image_count, restore_portrait=not next_wide)
            i += 1
            continue
        if line.startswith('```'):
            lang=line[3:].strip();j=i+1
            while j<len(lines) and not lines[j].strip().startswith('```'):j+=1
            if lang=='mermaid':
                if fig>=len(figures):raise ValueError('Missing existing diagram asset')
                figure(doc,figures[fig],fig+1,root);fig+=1
            else:
                for code in lines[i+1:j]:
                    para=doc.add_paragraph();para.paragraph_format.space_after=Pt(1)
                    r=para.add_run(code);font(r,9.5);r.font.name='Consolas'
            i=j+1;continue
        if line.startswith('|') and i+1<len(lines) and re.match(r'^\|[ :|\-]+\|$',lines[i+1].strip()):
            j=i+2
            while j<len(lines) and lines[j].strip().startswith('|'):j+=1
            table(doc,lines[i:j],source,root,cutoff);i=j;continue
        match=re.match(r'^(#{2,4})\s+(.*)$',line)
        if match:
            # Title/heading punctuation is removed only in exported presentation, never in source meaning.
            heading=plain(match[2]);heading=re.sub(r'[^\w\u3400-\u9fff ]',' ',heading);heading=re.sub(r' +',' ',heading).strip()
            doc.add_heading(heading,len(match[1])-1);i+=1;continue
        if line.startswith('> '):
            para=paragraph(doc,line[2:],source,root,cutoff);para.paragraph_format.left_indent=Inches(.2)
        elif re.match(r'^[-*] ',line):paragraph(doc,line[2:],source,root,cutoff,style='List Bullet')
        elif re.match(r'^\d+\. ',line):paragraph(doc,re.sub(r'^\d+\. ','',line),source,root,cutoff,style='List Number')
        else:paragraph(doc,line,source,root,cutoff)
        i+=1
    if fig and fig!=len(figures):raise ValueError('Existing diagram count differs from Markdown')
    return image_count + fig


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root',type=Path,required=True)
    parser.add_argument('--private-root',type=Path,required=True)
    parser.add_argument('--cutoff',required=True)
    parser.add_argument('--documentation-ref',default='HEAD')
    parser.add_argument('--stage', choices=['M6', 'M7', 'FINAL-20261029', 'FINAL90b'], required=True)
    parser.add_argument('--only', nargs='+', choices=[row[0] for row in REPORTS])
    args=parser.parse_args();root=args.root.resolve();private=args.private_root.resolve()
    args.cutoff = subprocess.check_output(['git','rev-parse','--verify',args.cutoff+'^{commit}'],cwd=root,text=True).strip()
    args.documentation_ref = subprocess.check_output(['git','rev-parse','--verify',args.documentation_ref+'^{commit}'],cwd=root,text=True).strip()
    if not str(Path(sys.executable)).startswith(str(RUNTIME/'python')):
        raise RuntimeError('Use the resolved bundled Python')
    if private.name!='.submission-private':raise ValueError('Private output root required')
    cover_file=private/'cover-info.json'
    if stat.S_IMODE(cover_file.stat().st_mode)!=0o600:raise ValueError('Private cover must have mode0600')
    info=json.loads(cover_file.read_text());os.umask(0o077)
    if args.stage == 'M7':
        if args.only and args.only != ['test-results-report.md']:
            raise ValueError('M7 only authors the substantively updated test report')
        args.only = ['test-results-report.md']
    if not (private/'review-documents'/'M4'/'archive-integrity.json').exists():
        raise ValueError('Archive the original M4 files byte-identically before M6 authoring')
    out=private/'review-documents'/args.stage;out.mkdir(parents=True,exist_ok=True);out.chmod(0o700)
    reused = []
    if args.stage == 'M7':
        previous = private/'review-documents'/'M6'
        for row in json.loads((previous/'export-manifest.json').read_text())['reports']:
            if row['source'] == 'reports/final/test-results-report.md': continue
            if hashlib.sha256((root/row['source']).read_bytes()).hexdigest() != row['source_sha256']:
                raise ValueError('An unchanged M6 report source differs; do not silently reexport it')
            if hashlib.sha256((previous/row['output']).read_bytes()).hexdigest() != row['output_sha256']:
                raise ValueError('An unchanged M6 Word hash differs')
            if not row['visual_QA'].startswith('passed'):
                raise ValueError('M6 reused Word must already have every-page QA')
            reused.append({**row, 'output': '../M6/'+row['output'], 'reexported': False})
        if len(reused) != 3: raise ValueError('M7 must reuse exactly three unchanged M6 documents')
    # Bundled LO uses a private fontconfig override; no global user font settings change.
    bundled_fonts = RUNTIME/'native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/Resources/fonts/truetype'
    config = ('<?xml version="1.0"?>\n<!DOCTYPE fontconfig SYSTEM "urn:fontconfig:fonts.dtd">\n'
              '<fontconfig>\n<dir>/System/Library/Fonts/Supplemental</dir>\n'
              '<dir>/System/Library/Fonts</dir>\n<dir>' + escape(str(bundled_fonts)) + '</dir>\n'
              '<cachedir>' + escape(str(out/'font-cache')) + '</cachedir>\n'
              '<alias><family>sans-serif</family><prefer><family>Songti SC</family></prefer></alias>\n'
              '<alias><family>serif</family><prefer><family>Songti SC</family></prefer></alias>\n</fontconfig>\n')
    (out/'fontconfig.xml').write_text(config)
    figures=json.loads((root/'docs/diagrams/manifest.json').read_text())['figures']
    for item in figures:
        ASSET_HASHES[item['artifacts']['png']['path']] = item['artifacts']['png']['sha256']
    for name in ['M5-render-manifest.json', 'M6-render-manifest.json', 'FINAL-20261029-render-manifest.json']:
        for item in json.loads((root/'docs/reports/diagrams'/name).read_text())['artifacts']:
            if item['output'].endswith('.png'): ASSET_HASHES[item['output']] = item['sha256']
    prior_path = out/'export-manifest.json'
    records = json.loads(prior_path.read_text())['reports'] if prior_path.exists() else []
    exporter_hash = hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
    for name,target,title in REPORTS:
        if args.stage in {'FINAL-20261029', 'FINAL90b'}: target = name.replace('.md', '.docx')
        if args.only and name not in args.only: continue
        source=root/'reports/final'/name;doc=Document();configure(doc,title);cover(doc,title,info,args.cutoff)
        selected=[f for f in figures if f['source']==source.relative_to(root).as_posix()]
        source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        prior = next((row for row in records if row['source'] == source.relative_to(root).as_posix()), None)
        if prior and prior['source_sha256'] == source_hash and prior.get('exporter_sha256') == exporter_hash and prior.get('documentation_ref') == args.documentation_ref and (out/target).exists():
            continue
        figure_count = body(doc,source,root,args.documentation_ref,selected)
        output=out/target
        if output.exists():
            archive = out/'iterations'/datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')
            archive.mkdir(parents=True, mode=0o700)
            shutil.copy2(output, archive/target)
            if prior_path.exists(): shutil.copy2(prior_path, archive/'export-manifest-before.json')
        doc.save(output);output.chmod(0o600)
        records = [row for row in records if row['source'] != source.relative_to(root).as_posix()]
        records.append({'source':source.relative_to(root).as_posix(),'source_sha256':source_hash,
                        'output':target,'output_sha256':hashlib.sha256(output.read_bytes()).hexdigest(),
                        'figures':figure_count,'cutoff':args.cutoff,'documentation_ref':args.documentation_ref,'date_blank':True,
                        'exporter_sha256':exporter_hash,'visual_QA':'pending rendering and every-page inspection'})
    (out/'export-manifest.json').write_text(json.dumps({'reports':records,'private_cover_used':True,
        'stage':args.stage,'M6_reused_reports':reused,'M4_archive_preserved':True,
        'operation_marker': ('new create4 marker required once immediately before first authoring' if args.stage in {'FINAL-20261029', 'FINAL90b'} else 'original M4 operation not repeated'),
        'student_research_authored':False,'visual_QA':'pending rendering and every-page inspection'},ensure_ascii=False,indent=2)+'\n')
    print('Selected private technical review DOCX sources exported incrementally; render and every-page QA still required.')


if __name__=='__main__':main()
